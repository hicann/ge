#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# -----------------------------------------------------------------------------------------------------------
# Copyright (c) 2026 Huawei Technologies Co., Ltd.
# This program is free software, you can redistribute it and/or modify it under the terms and conditions of
# CANN Open Software License Agreement Version 2.0 (the "License").
# Please refer to the License for details. You may not use this file except in compliance with the License.
# THIS SOFTWARE IS PROVIDED ON AN "AS IS" BASIS, WITHOUT WARRANTIES OF ANY KIND, EITHER EXPRESS OR IMPLIED,
# INCLUDING BUT NOT LIMITED TO NON-INFRINGEMENT, MERCHANTABILITY, OR FITNESS FOR A PARTICULAR PURPOSE.
# See LICENSE in the root of the software repository for the full text of the License.
# -----------------------------------------------------------------------------------------------------------

"""Generate Excel report from PR log check findings.

Usage:
    python3 gen_report.py --findings findings.json --output report.xlsx

findings.json format: a JSON array of objects with keys:
    severity, file, line, rule, description, suggestion, code_snippet
"""

import argparse
import json
import logging
import sys
from collections import Counter
from pathlib import Path

logging.basicConfig(
    level=logging.INFO, format="[%(levelname)s] %(asctime)s : %(message)s"
)

SEVERITY_ORDER = {"致命": 0, "严重": 1, "中等": 2, "提示": 3}
SEVERITY_COLORS = {
    "致命": "FF5050",
    "严重": "F8CBAD",
    "中等": "FFE699",
    "提示": "E2EFDA",
}
SEVERITY_TEXT_COLORS = {
    "致命": "FFFFFF",
    "严重": "000000",
    "中等": "000000",
    "提示": "000000",
}


def gen_report(findings, output_path, meta=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        logging.error("openpyxl not installed, run: pip install openpyxl")
        return 1

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "概览"
    sev_counter = Counter(f["severity"] for f in findings)
    file_set = {f["file"] for f in findings}

    overview_rows = [
        ("检查项", "值"),
        ("发现文件数", len(file_set)),
        ("问题总数", len(findings)),
        ("致命", sev_counter.get("致命", 0)),
        ("严重", sev_counter.get("严重", 0)),
        ("中等", sev_counter.get("中等", 0)),
        ("提示", sev_counter.get("提示", 0)),
    ]
    if meta:
        for k, v in meta.items():
            overview_rows.append((k, v))

    for row_idx, (key, val) in enumerate(overview_rows, 1):
        ws1.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        ws1.cell(row=row_idx, column=2, value=val)
    ws1.column_dimensions["A"].width = 20
    ws1.column_dimensions["B"].width = 40

    ws2 = wb.create_sheet("问题明细")
    headers = [
        "序号",
        "严重级",
        "文件",
        "行号",
        "规则",
        "问题描述",
        "修复建议",
        "代码片段",
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sorted_findings = sorted(
        findings,
        key=lambda f: (
            SEVERITY_ORDER.get(f["severity"], 99),
            f.get("file", ""),
            f.get("line", 0),
        ),
    )
    for row_idx, f in enumerate(sorted_findings, 2):
        severity = f.get("severity", "")
        color = SEVERITY_COLORS.get(severity)
        text_color = SEVERITY_TEXT_COLORS.get(severity, "000000")
        font = Font(color=text_color, bold=(severity == "致命"))
        fill = (
            PatternFill(start_color=color, end_color=color, fill_type="solid")
            if color
            else None
        )

        values = [
            row_idx - 1,
            severity,
            f.get("file", ""),
            f.get("line", ""),
            f.get("rule", ""),
            f.get("description", ""),
            f.get("suggestion", ""),
            f.get("code_snippet", ""),
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill
            cell.font = font

    col_widths = [6, 8, 45, 8, 12, 50, 45, 60]
    for col_idx, width in enumerate(col_widths, 1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = width
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(sorted_findings) + 1}"
    )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    logging.info("Report saved to %s", output)
    logging.info("  Total findings: %d", len(findings))
    logging.info(
        "  Fatal: %s, Severe: %s, Medium: %s, Info: %s",
        sev_counter.get("致命", 0),
        sev_counter.get("严重", 0),
        sev_counter.get("中等", 0),
        sev_counter.get("提示", 0),
    )
    return 0


def main():
    parser = argparse.ArgumentParser(
        description="Generate Excel report from PR log check findings"
    )
    parser.add_argument("--findings", required=True, help="Path to findings JSON file")
    parser.add_argument(
        "--output", default="pr-log-check-report.xlsx", help="Output Excel path"
    )
    parser.add_argument(
        "--meta", help="Optional JSON file with extra overview metadata"
    )
    args = parser.parse_args()

    findings_path = Path(args.findings)
    if not findings_path.exists():
        logging.error("Error: findings file not found: %s", findings_path)
        return 1
    with open(findings_path, "r", encoding="utf-8") as fh:
        findings = json.load(fh)
    if not isinstance(findings, list):
        logging.error(
            "Error: findings must be a JSON array, got %s", type(findings).__name__
        )
        return 1

    meta = None
    if args.meta:
        with open(args.meta, "r", encoding="utf-8") as fh:
            meta = json.load(fh)

    return gen_report(findings, args.output, meta)


if __name__ == "__main__":
    sys.exit(main())
